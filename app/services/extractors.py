"""Format-specific text extraction kept separate from processing."""

from pathlib import Path
from datetime import date, datetime, time
from typing import Protocol


class ExtractionError(ValueError):
    """Raised when a supported file cannot be converted to text."""


class TextExtractor(Protocol):
    """Interface implemented by document-format extractors."""

    def extract(self, path: Path) -> str: ...


class TxtExtractor:
    def extract(self, path: Path) -> str:
        try:
            return path.read_text(encoding="utf-8")
        except UnicodeDecodeError as exc:
            raise ExtractionError(
                f"Text document must use UTF-8 encoding: {path.name}"
            ) from exc


class PdfExtractor:
    def extract(self, path: Path) -> str:
        try:
            from pypdf import PdfReader

            reader = PdfReader(path)
            return "\n\n".join(page.extract_text() or "" for page in reader.pages)
        except ImportError as exc:
            raise ExtractionError("PDF support requires the pypdf package") from exc
        except Exception as exc:
            raise ExtractionError(f"Could not extract PDF text: {path.name}") from exc


class DocxExtractor:
    def extract(self, path: Path) -> str:
        try:
            from docx import Document

            document = Document(path)
            sections = [
                paragraph.text.strip()
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]
            for table_index, table in enumerate(document.tables, start=1):
                rows = [
                    [cell.text.replace("\n", " ").strip() for cell in row.cells]
                    for row in table.rows
                ]
                if not rows:
                    continue
                headers = [
                    value or f"Column {index + 1}"
                    for index, value in enumerate(rows[0])
                ]
                table_lines = [f"Table {table_index}"]
                for row_number, values in enumerate(rows[1:], start=2):
                    pairs = [
                        f"{headers[index]}={value}"
                        for index, value in enumerate(values)
                        if value
                    ]
                    if pairs:
                        table_lines.append(f"Row {row_number}: " + " | ".join(pairs))
                if len(table_lines) > 1:
                    sections.append("\n".join(table_lines))
            return "\n\n".join(sections)
        except ImportError as exc:
            raise ExtractionError("DOCX support requires the python-docx package") from exc
        except Exception as exc:
            raise ExtractionError(f"Could not extract DOCX text: {path.name}") from exc


class XlsxExtractor:
    """Convert workbook structure into searchable, header-aware text."""

    max_sheets = 50
    max_rows_per_sheet = 100_000
    max_columns = 200

    def extract(self, path: Path) -> str:
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(
                path, read_only=True, data_only=False, keep_links=False
            )
            values_workbook = load_workbook(
                path, read_only=True, data_only=True, keep_links=False
            )
            sections: list[str] = []
            for sheet_index, sheet in enumerate(workbook.worksheets):
                if sheet_index >= self.max_sheets:
                    sections.append(
                        f"Workbook note: remaining sheets omitted after {self.max_sheets}."
                    )
                    break
                value_sheet = values_workbook[sheet.title]
                section = self._sheet_text(sheet, value_sheet)
                if section:
                    sections.append(section)
            workbook.close()
            values_workbook.close()
            return "\n\n".join(sections)
        except ImportError as exc:
            raise ExtractionError("Excel support requires the openpyxl package") from exc
        except Exception as exc:
            raise ExtractionError(f"Could not extract Excel workbook: {path.name}") from exc

    def _sheet_text(self, sheet, value_sheet) -> str:
        lines = [f"Worksheet: {sheet.title}"]
        headers: list[str] | None = None
        rows_seen = 0
        paired_rows = zip(
            sheet.iter_rows(max_col=self.max_columns),
            value_sheet.iter_rows(max_col=self.max_columns),
            strict=False,
        )
        for row_number, (formula_row, value_row) in enumerate(paired_rows, start=1):
            if row_number > self.max_rows_per_sheet:
                lines.append(
                    f"Worksheet note: remaining rows omitted after {self.max_rows_per_sheet}."
                )
                break
            values = [
                self._cell_text(formula_cell.value, value_cell.value)
                for formula_cell, value_cell in zip(
                    formula_row, value_row, strict=False
                )
            ]
            while values and not values[-1]:
                values.pop()
            if not any(values):
                continue
            rows_seen += 1
            if headers is None:
                headers = [value or f"Column {index + 1}" for index, value in enumerate(values)]
                pairs = [f"{self._column_name(index)}={value}" for index, value in enumerate(values) if value]
            else:
                pairs = []
                for index, value in enumerate(values):
                    if not value:
                        continue
                    label = headers[index] if index < len(headers) else self._column_name(index)
                    pairs.append(f"{label}={value}")
            lines.append(f"Row {row_number}: " + " | ".join(pairs))
        return "\n".join(lines) if rows_seen else ""

    @staticmethod
    def _cell_text(formula_value, calculated_value) -> str:
        value = calculated_value if calculated_value is not None else formula_value
        if value is None:
            return ""
        if isinstance(value, (datetime, date, time)):
            rendered = value.isoformat()
        elif isinstance(value, bool):
            rendered = "true" if value else "false"
        else:
            rendered = str(value).replace("\r", " ").replace("\n", " ").strip()
        if (
            isinstance(formula_value, str)
            and formula_value.startswith("=")
        ):
            if calculated_value is not None:
                return f"{rendered} (formula: {formula_value})"
            return f"formula: {formula_value}"
        return rendered

    @staticmethod
    def _column_name(index: int) -> str:
        value = index + 1
        name = ""
        while value:
            value, remainder = divmod(value - 1, 26)
            name = chr(65 + remainder) + name
        return f"Column {name}"
