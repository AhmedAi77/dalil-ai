from pathlib import Path

from docx import Document

from app.services.document_processor import DocumentProcessor


def test_docx_uses_same_structured_output(tmp_path: Path) -> None:
    path = tmp_path / "guide.docx"
    source = Document()
    source.add_paragraph("Docker guide")
    source.add_paragraph("Restart the daemon after checking logs.")
    source.save(path)

    result = DocumentProcessor().process(path)

    assert result.extension == ".docx"
    assert "Docker guide" in result.content
    assert "Restart the daemon" in result.content


def test_docx_extracts_table_rows_with_headers(tmp_path: Path) -> None:
    path = tmp_path / "support-policy.docx"
    source = Document()
    source.add_paragraph("Cedar Support Policy")
    table = source.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Policy item"
    table.rows[0].cells[1].text = "Requirement"
    cells = table.add_row().cells
    cells[0].text = "Priority Red response"
    cells[1].text = "Within 15 minutes"
    source.save(path)

    result = DocumentProcessor().process(path)

    assert "Table 1" in result.content
    assert "Policy item=Priority Red response" in result.content
    assert "Requirement=Within 15 minutes" in result.content
